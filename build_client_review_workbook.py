#!/usr/bin/env python3

from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


INPUT_FILE = "telos_matches_enriched_v3.csv"
OUTPUT_FILE = "telos_matches_client_review_workbook.xlsx"

CLIENT_INPUT_COLUMNS = [
    "client_confirmed_match",
    "client_confirmed_country",
    "client_confirmed_address",
    "client_confirmed_postcode",
    "client_notes",
]

COLUMN_RENAMES = {
    "supplier_legal_name": "client_supplier_name",
    "_supplier_match_name": "supplier_match_name_internal",
    "entity_name": "matched_register_name",
    "matched_entity_name": "matched_name_norm_internal",
    "matched_register": "matched_register_id",
    "ei_registration_number": "master_register_id",
}


def load_data(path: str) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str, keep_default_na=False)

    for col in CLIENT_INPUT_COLUMNS:
        if col not in df.columns:
            df[col] = ""

    df = df.rename(columns=COLUMN_RENAMES)

    preferred = [
        "client_supplier_name",
        "supplier_match_name_internal",
        "match_type",
        "match_score",
        "confidence",
        "matched_register_name",
        "matched_name_norm_internal",
        "matched_register_id",
        "master_register_id",
        "country",
        "ccaa",
        "city",
        "postcode",
        "address",
        "tax_id",
        "client_confirmed_match",
        "client_confirmed_country",
        "client_confirmed_address",
        "client_confirmed_postcode",
        "client_notes",
    ]

    ordered = [c for c in preferred if c in df.columns] + [c for c in df.columns if c not in preferred]
    df = df[ordered]

    return df


def apply_header_style(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_column_widths(ws):
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = 0
        col_name = ws.cell(row=1, column=col_idx).value or ""

        for cell in col_cells[:300]:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(val))

        if col_name in {"address", "client_notes"}:
            width = 40
        elif col_name in {"client_supplier_name", "matched_register_name", "matched_name_norm_internal"}:
            width = 28
        elif col_name in {"supplier_match_name_internal"}:
            width = 24
        else:
            width = min(max(max_len + 2, 12), 24)

        ws.column_dimensions[get_column_letter(col_idx)].width = width


def add_confidence_formatting(ws, headers):
    if "confidence" not in headers:
        return

    conf_col = headers.index("confidence") + 1

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=conf_col)
        value = (cell.value or "").strip().upper()

        if value == "HIGH":
            cell.fill = PatternFill("solid", fgColor="C6E0B4")
        elif value == "MEDIUM":
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
        elif value == "LOW":
            cell.fill = PatternFill("solid", fgColor="F4CCCC")


def unlock_client_input_columns(ws, headers):
    locked = Protection(locked=True)
    unlocked = Protection(locked=False)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.protection = locked

    for col_name in CLIENT_INPUT_COLUMNS:
        if col_name in headers:
            col_idx = headers.index(col_name) + 1
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col_idx).protection = unlocked


def add_validation(ws, headers):
    if "client_confirmed_match" in headers:
        col_idx = headers.index("client_confirmed_match") + 1
        col_letter = get_column_letter(col_idx)

        dv = DataValidation(
            type="list",
            formula1='"YES,NO,UNSURE"',
            allow_blank=True
        )
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}{ws.max_row}")


def protect_sheet(ws):
    ws.protection.sheet = True
    ws.protection.autoFilter = True
    ws.protection.sort = True
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def write_dataframe_sheet(wb, sheet_name: str, df: pd.DataFrame):
    ws = wb.create_sheet(title=sheet_name)

    headers = list(df.columns)
    ws.append(headers)

    for row in df.itertuples(index=False):
        ws.append(list(row))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    apply_header_style(ws)
    set_column_widths(ws)
    add_confidence_formatting(ws, headers)
    unlock_client_input_columns(ws, headers)
    add_validation(ws, headers)
    protect_sheet(ws)

    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=1).alignment = Alignment(vertical="top")
        if "client_notes" in headers:
            notes_col = headers.index("client_notes") + 1
            ws.cell(row=row, column=notes_col).alignment = Alignment(wrap_text=True, vertical="top")


def add_instructions_sheet(wb):
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_view.showGridLines = False

    rows = [
        ["Telos supplier match review"],
        [""],
        ["What this workbook is"],
        ["This workbook shows client supplier names alongside suggested register matches and available enrichment fields."],
        [""],
        ["How to review"],
        ["1. Start with Review_Exact, then Review_Fuzzy."],
        ["2. Confirm whether the matched register name is the same supplier your team intended."],
        ["3. Use country, city, postcode, and address to triangulate the correct entity."],
        ["4. Fill in only the editable columns."],
        [""],
        ["Editable columns"],
        ["client_confirmed_match → YES / NO / UNSURE"],
        ["client_confirmed_country → enter the country if known"],
        ["client_confirmed_address → enter/correct address if known"],
        ["client_confirmed_postcode → enter/correct postcode if known"],
        ["client_notes → any context, caveats, or corrections"],
        [""],
        ["Locked columns"],
        ["All system-generated columns are locked to preserve auditability."],
        [""],
        ["Important"],
        ["client_supplier_name is the original supplier name from the client file and should be used for reconciliation."],
        ["matched_register_name is the matched name from the register/master data."],
    ]

    for r in rows:
        ws.append(r)

    ws["A1"].font = Font(size=14, bold=True)
    ws["A3"].font = Font(bold=True)
    ws["A6"].font = Font(bold=True)
    ws["A12"].font = Font(bold=True)
    ws["A19"].font = Font(bold=True)
    ws["A22"].font = Font(bold=True)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 110
    ws.protection.sheet = True


def main():
    input_path = Path(INPUT_FILE)
    output_path = Path(OUTPUT_FILE)

    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    df = load_data(str(input_path))

    wb = Workbook()
    add_instructions_sheet(wb)

    write_dataframe_sheet(wb, "Review_All", df)

    exact = df[df["match_type"] == "name_exact_norm"].copy() if "match_type" in df.columns else df.iloc[0:0].copy()
    fuzzy = df[df["match_type"] == "name_fuzzy"].copy() if "match_type" in df.columns else df.iloc[0:0].copy()

    write_dataframe_sheet(wb, "Review_Exact", exact)
    write_dataframe_sheet(wb, "Review_Fuzzy", fuzzy)

    wb.save(output_path)

    print(f"Saved: {output_path}")
    print(f"All rows: {len(df)}")
    print(f"Exact rows: {len(exact)}")
    print(f"Fuzzy rows: {len(fuzzy)}")


if __name__ == "__main__":
    main()
